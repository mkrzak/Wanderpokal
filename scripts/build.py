#!/usr/bin/env python3
"""Build the Wanderpokal site from the match spreadsheet.

Usage:
    python3 scripts/build.py path/to/Wanderpokal.xlsx

Reads every weekly match sheet ("89 - 13102026"), every special night
("Bryan Cup - 2810205", ...) and the "WinnersList & Stats" sheet, then writes:
    data/wanderpokal.json   – all nights, rosters and set scores
    index.html              – the site, with the data embedded (works offline
                              and on GitHub Pages)

Requires: pip install openpyxl
"""
import json, re, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is missing: run  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent

# Sheets that are not Tuesday match nights.
SKIP = {
    'Tuesday Match Template', 'Registration 2025', 'Registration 2026', 'WinnersList & Stats',
    'xmas Rules', 'xmas_results matrix', 'RULES (SL26)', 'comments and payments etc',
    'Fixtures (SL26) - 18042026', 'Results Matrix (SL26) - Group S', 'Results Matrix (SL26)- FINALS',
    'RULES (BD26)', 'Fixtures (BD26)', 'monikas birthday', 'CRS reunion - 05052026',
}
# Spelling fixes so one person keeps one row in the leaderboard. Add more here.
ALIASES = {
    'Federica C': 'Fede C', 'Vagelis': 'Vangelis', 'allen': 'Allen', 'Attturk': 'Ataturk',
    'Emanuala': 'Emanuela', 'MIsh': 'Mish', 'MItesh': 'Mitesh', 'Niocole': 'Nicole', 'Veronika': 'Veronica', 'James': 'James Y', 'Ceejay': 'CJ', 'Alex Q': 'Alex',
}
# Players swapped between teams on a weekly match sheet: week -> [(name, name), ...]
SWAPS = {17: [('Jack', 'Mish')], 20: [('Karolina', 'Fede B')]}
# Dates the sheet names don't carry.
DATE_FIX = {'Banana vs Wanderpokal Pt.1 - 14': '2026-07-14', 'Banana vs Wanderpokal Pt.2 - 21': '2026-07-21'}


def name(x):
    x = str(x).strip()
    return ALIASES.get(x, x)


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def parse_night(ws):
    grid = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    title = ws.title.strip()

    date = DATE_FIX.get(title)
    m = re.search(r'(\d{2})(\d{2})(\d{3,4})\s*$', title)
    if not date and m:
        d, mo, y = m.groups()
        if len(y) == 3:
            y = '2025'                                       # "Bryan Cup - 2810205" typo
        date = f'{y}-{mo}-{d}'
    wk = re.match(r'^\s*(\d+)\s*-\s*\d{8}\s*$', title)  # "89 - 13102026" but not "5-1 Special - ..."
    week = int(wk.group(1)) if wk else None
    label = None if week else re.sub(r'\s*-\s*\d+\s*$', '', title).strip()

    # Rosters: columns under the "TEAM A | TEAM B | TEAM C" header row.
    teams = {'A': [], 'B': [], 'C': []}
    anchor = next(((r, c) for (r, c), v in grid.items()
                   if isinstance(v, str) and v.strip().upper().startswith('TEAM A')
                   and str(grid.get((r, c + 1), '')).strip().upper().startswith('TEAM B')), None)
    if anchor:
        r0, c0 = anchor
        for i, t in enumerate('ABC'):
            for r in range(r0 + 1, r0 + 12):
                if 'Total' in str(grid.get((r, c0 + 3), '')) or 'Total' in str(grid.get((r, c0 - 1), '')):
                    break
                v = grid.get((r, c0 + i))
                if isinstance(v, str) and v.strip() and v.strip().upper() not in ('1ST', '2ND', '3RD') \
                        and 'winner' not in v.lower() and ' won ' not in v.lower():
                    teams[t].append(name(v))

    # Sets: a row "A | B | 20 | 18".
    matches = []
    for (r, c), v in sorted(grid.items()):
        if isinstance(v, str) and v.strip() in ('A', 'B', 'C'):
            v2 = grid.get((r, c + 1))
            if isinstance(v2, str) and v2.strip() in ('A', 'B', 'C') and v2.strip() != v.strip():
                s1, s2 = grid.get((r, c + 2)), grid.get((r, c + 3))
                if is_num(s1) and is_num(s2):
                    matches.append([v.strip(), v2.strip(), int(s1), int(s2)])

    if not matches and not any(teams.values()):
        return None
    return {'id': re.sub(r'[^A-Za-z0-9]+', '-', title).strip('-').lower(),
            'week': week, 'name': label, 'date': date, 'teams': teams, 'matches': matches}


PHOTO_NAME = re.compile(r'^(\d{4}-\d{2}-\d{2})_(?:week(\d+)|special)(?:_\d+)?\.(?:jpe?g|png)$', re.I)


def add_new_photos(photos_file, history, special_winners):
    """Add photos dropped into photos/ (e.g. 2026-10-27_week90.jpg, 2026-11-03_special_2.jpg)
    to data/photos.json and make their thumbnails. Photos uploaded from the website are already listed."""
    folder = ROOT / 'photos'
    data = json.loads(photos_file.read_text()) if photos_file.exists() else {'cards': []}
    known = {ph['f'] for c in data['cards'] for ph in c['photos']}
    new = sorted(f for f in folder.glob('*') if f.is_file() and PHOTO_NAME.match(f.name) and f.name not in known) if folder.exists() else []
    if not new:
        return
    try:
        from PIL import Image, ImageOps
    except ImportError:
        print('New photos found but Pillow is missing (pip install pillow): skipped', ', '.join(f.name for f in new))
        return
    names = {h['date']: h.get('name') for h in history if not h['week']}
    names.update({sw['date']: sw['name'] for sw in special_winners})
    (folder / 'thumbs').mkdir(exist_ok=True)
    for f in new:
        date, week = PHOTO_NAME.match(f.name).groups()
        with Image.open(f) as im:
            im = ImageOps.exif_transpose(im).convert('RGB')
            if f.suffix.lower() != '.jpg' or max(im.size) > 1600:      # keep the site light: 1600px JPG
                im.thumbnail((1600, 1600))
                f.unlink(); f = f.with_suffix('.jpg'); im.save(f, quality=85)
            w, h = im.size
            th = im.copy(); th.thumbnail((720, 99999)); th.save(folder / 'thumbs' / f.name, quality=80)
        card = next((c for c in data['cards'] if not c.get('approx') and c['date'] == date), None)
        if not card:
            card = {'kind': 'weekly', 'date': date, 'week': int(week), 'photos': []} if week else \
                   {'kind': 'special', 'date': date, 'name': names.get(date) or 'Special', 'photos': []}
            data['cards'].append(card)
        card['photos'].append({'f': f.name, 'w': w, 'h': h})
    data['cards'].sort(key=lambda c: c.get('date') or '', reverse=True)
    photos_file.write_text(json.dumps(data, indent=1))
    print(f'Added {len(new)} new photo(s):', ', '.join(f.name for f in new))



# ---------- private gender statistics (totals only go into the page) ----------
def _g_solve(A, b):
    n = len(b); M = [row[:] + [b[i]] for i, row in enumerate(A)]
    for c in range(n):
        p = max(range(c, n), key=lambda r: abs(M[r][c])); M[c], M[p] = M[p], M[c]
        for r in range(c + 1, n):
            f = M[r][c] / M[c][c]
            if f:
                Mr, Mc = M[r], M[c]
                for k in range(c, n + 1):
                    Mr[k] -= f * Mc[k]
    x = [0.0] * n
    for r in range(n - 1, -1, -1):
        x[r] = (M[r][n] - sum(M[r][k] * x[k] for k in range(r + 1, n))) / M[r][r]
    return x


def _g_fit(rows, lam, players=True, exp=True):
    """Point difference of a set ~ women + men (and others) on each side [+ team experience] [+ each player's skill,
    pulled towards 0 by lam]. Gender stays unpenalised, so it measures the gap between groups after skill."""
    pl = sorted({p for r in rows for p in r[2]}) if players else []
    fx = [0, 1] + ([2] if exp else [])
    idx = {p: len(fx) + i for i, p in enumerate(pl)}; n = len(fx) + len(pl)
    A = [[0.0] * n for _ in range(n)]; b = [0.0] * n
    for _, f, sides, y in rows:
        v = [(j, f[k]) for j, k in enumerate(fx) if f[k]] + ([(idx[p], s) for p, s in sides.items()] if players else [])
        for j, vj in v:
            b[j] += vj * y; Aj = A[j]
            for k, vk in v:
                Aj[k] += vj * vk
    for p in pl:
        A[idx[p]][idx[p]] += lam
    for j in range(len(fx)):
        A[j][j] += 1e-6
    x = _g_solve(A, b)
    return x[0], x[1]                      # per woman, per man/other


def gender_stats(history):
    """Totals only, from the private wanderpokal_players_gender.xlsx (never uploaded). No player's gender goes into the page."""
    import math, random
    f = ROOT / 'wanderpokal_players_gender.xlsx'
    if not f.exists():
        return None
    ws = openpyxl.load_workbook(f, data_only=True)['Players']
    rr = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rr) if r and 'Player' in r)
    head = list(rr[hi]); pc, gc = head.index('Player'), head.index('Gender')
    gender = {name(r[pc]): str(r[gc] or '').strip().lower() for r in rr[hi + 1:] if r and r[pc] and r[gc]}
    woman = {p for p, g in gender.items() if g == 'woman'}
    rows, seen, ni, groups, per = [], {}, 0, {}, {}
    for h in sorted(history, key=lambda h: h['date'] or ''):
        m = h.get('matches') or []
        if not h['week'] or len(m) != 9 or any(x[2] is None or x[3] is None for x in m):
            continue
        if any(p not in gender for t in h['teams'].values() for p in t):
            continue
        T = {t: list(dict.fromkeys(h['teams'][t])) for t in 'ABC'}
        ex = {p: math.log1p(seen.get(p, 0)) for t in T.values() for p in t}
        feat = {t: (sum(p in woman for p in T[t]), sum(p not in woman for p in T[t]), sum(ex[p] for p in T[t]) / len(T[t])) for t in 'ABC'}
        W = {t: [0, 0, 0, 0] for t in 'ABC'}            # set wins, points for, points against, sets
        for x, y, a, b in m:
            fx, fy = feat[x], feat[y]
            sides = {p: 1.0 for p in T[x]}
            for p in T[y]:
                sides[p] = sides.get(p, 0) - 1.0
            rows.append((ni, [fx[i] - fy[i] for i in range(3)], sides, a - b))
            for team, other, pf, pa in ((x, y, a, b), (y, x, b, a)):
                d = max(-2, min(2, feat[team][0] - feat[other][0]))
                groups.setdefault(d, []).append((pf - pa, 1 if pf > pa else 0))
                W[team][0] += pf > pa; W[team][1] += pf; W[team][2] += pa; W[team][3] += 1
        order = sorted('ABC', key=lambda t: (-W[t][0], -W[t][1]))
        for t in 'ABC':
            for p in T[t]:
                q = per.setdefault(p, [0, 0, 0, 0, 0])   # nights, titles, set wins, sets, points for
                q[0] += 1; q[1] += order[0] == t; q[2] += W[t][0]; q[3] += W[t][3]; q[4] += W[t][1]
        for t in T.values():
            for p in t:
                seen[p] = seen.get(p, 0) + 1
        ni += 1
    if ni < 10:
        return None
    by = {}
    for r in rows:
        by.setdefault(r[0], []).append(r)
    models = []
    rnd = random.Random(2024)
    for label, kw in (('Simple: women and men counted', dict(players=False, exp=False)),
                      ('+ experience', dict(players=False, exp=True)),
                      ('+ each player’s skill (full model)', dict(players=True, exp=True))):
        w, o = _g_fit(rows, 30, **kw); boots = []
        for _ in range(200):                     # resample whole nights
            smp = [r for n in rnd.choices(range(ni), k=ni) for r in by[n]]
            bw, bo = _g_fit(smp, 30, **kw); boots.append(bw - bo)
        boots.sort()
        models.append({'label': label, 'est': round(w - o, 2), 'lo': round(boots[4], 2), 'hi': round(boots[194], 2)})
    bars = []
    for d in sorted(groups):
        v = groups[d]; k = len(v); md = sum(e[0] for e in v) / k
        sd = (sum((e[0] - md) ** 2 for e in v) / (k - 1)) ** 0.5 if k > 1 else 0
        bars.append({'diff': d, 'n': k, 'pts': round(md, 2), 'ci': round(1.96 * sd / k ** 0.5, 2), 'win': round(sum(e[1] for e in v) / k, 3)})
    def summary(g):                                  # averages over players with 10+ nights; no headcounts published
        ps = [q for p, q in per.items() if gender.get(p) == g and q[0] >= 10]
        if len(ps) < 5:
            return None
        avg = lambda f: round(sum(f(q) for q in ps) / len(ps), 3)
        return {'nights': avg(lambda q: q[0]), 'nightWin': avg(lambda q: q[1] / q[0]), 'setWin': avg(lambda q: q[2] / q[3]), 'ppm': avg(lambda q: q[4] / q[3])}
    return {'nights': ni, 'sets': len(rows), 'groups': bars, 'models': models, 'women': summary('woman'), 'men': summary('man')}


def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # Winners list: covers early weeks that have no match sheet.
    winners = {}
    ws = wb['WinnersList & Stats']
    for r in range(3, ws.max_row + 1):
        wkv, team = ws.cell(r, 1).value, ws.cell(r, 2).value
        m = re.search(r'wk\s*0*(\d+)', str(wkv or ''))
        if not m:
            continue
        dm = re.match(r'(\d\d)_(\d\d)_(\d\d)', str(wkv))
        winners[int(m.group(1))] = {
            'team': str(team or '').replace('Team', '').strip(),
            'players': [name(ws.cell(r, c).value) for c in range(3, 9) if ws.cell(r, c).value],
            'date': f'20{dm.group(1)}-{dm.group(2)}-{dm.group(3)}' if dm else None,
        }

    # Special events list (columns AV:BC of the same sheet): name, date, winning players.
    special_winners = []
    for r in range(3, ws.max_row + 1):
        ev, when = ws.cell(r, 48).value, ws.cell(r, 49).value
        if not ev or not hasattr(when, 'strftime'):
            continue
        special_winners.append({'name': str(ev).strip(), 'date': when.strftime('%Y-%m-%d'),
                                'players': [name(ws.cell(r, c).value) for c in range(50, 56) if ws.cell(r, c).value]})

    history, skipped = [], []
    for ws in wb:
        if ws.title in SKIP:
            continue
        night = parse_night(ws)
        if not night:
            skipped.append(ws.title)
            continue
        if not night['matches'] and night['week'] in winners:
            night['winner'] = winners[night['week']]['team']
        history.append(night)
    history.sort(key=lambda h: h['date'] or '')
    # Team corrections confirmed by Monika (the weekly match sheets still have these two players swapped)
    for h in history:
        for a, b in SWAPS.get(h['week'], []):
            for t, team in h['teams'].items():
                h['teams'][t] = [b if x == a else a if x == b else x for x in team]

    covered = {h['week'] for h in history if h['week']}
    legacy = [{'week': k, **v} for k, v in sorted(winners.items()) if k not in covered]
    for l in legacy:          # week 8 carries a mistyped date in the sheet
        if l['week'] == 8:
            l['date'] = None

    # Specials with a match sheet are already in history; keep the rest (winners only, or still to come).
    nice = {sw['date']: sw['name'] for sw in special_winners}
    for h in history:                      # use the tidy event names from the winners list
        if not h['week'] and h['date'] in nice:
            h['name'] = nice[h['date']]
    special_dates = {h['date'] for h in history if not h['week']}
    special_legacy = [sw for sw in special_winners if sw['date'] not in special_dates]

    data = {'history': history, 'legacy': legacy, 'specialLegacy': special_legacy}
    gs = gender_stats(history)
    if gs:
        data['genderStats'] = gs
    photos_file = ROOT / 'data' / 'photos.json'          # photo cards (see photos/ folder)
    add_new_photos(photos_file, history, special_winners)
    if photos_file.exists():
        data['photos'] = json.loads(photos_file.read_text())
        # left-to-right captions: only lines ticked "- [x]" in data/photo_captions.md
        cap_file = ROOT / 'data' / 'photo_captions.md'
        caps = {}
        if cap_file.exists():
            for line in cap_file.read_text().splitlines():
                m = re.match(r'^\s*-\s*\[[xX]\]\s*([^|]+?)\s*\|[^|]*\|\s*(.*)$', line)
                if m and m.group(2).strip():
                    caps[m.group(1).strip()] = [n.strip() for n in m.group(2).split(',') if n.strip()]
        for card in data['photos'].get('cards', []):
            for ph in card['photos']:
                ph.pop('ltr', None)
                if ph['f'] in caps:
                    ph['ltr'] = caps[ph['f']]
        print(f'{len(caps)} photo captions ticked')
    (ROOT / 'data').mkdir(exist_ok=True)
    (ROOT / 'data' / 'wanderpokal.json').write_text(json.dumps({k: v for k, v in data.items() if k != 'photos'}, indent=1))

    page = (ROOT / 'src' / 'template.html').read_text()
    embedded = json.dumps(data, separators=(',', ':')).replace('</', '<\\/')
    html = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
            '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">\n'
            '<meta name="description" content="Wanderpokal Tuesday volleyball: live scoring, player stats, results and records.">\n'
            '</head>\n<body>\n' + page.replace('__DATA__', embedded) + '\n</body>\n</html>\n')
    (ROOT / 'index.html').write_text(html)

    n_sp = sum(1 for h in history if not h['week'])
    print(f'{len(history) - n_sp} regular nights + {n_sp} specials with sheets, '
          f'{len(legacy)} winners-only weeks, {len(special_legacy)} winners-only/upcoming specials')
    if skipped:
        print('Skipped (no scores or rosters found):', ', '.join(skipped))
    print('Wrote index.html and data/wanderpokal.json')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
